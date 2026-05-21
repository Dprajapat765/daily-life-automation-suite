"""Excel output reporter — shared utility for writing styled Excel reports."""
from __future__ import annotations

from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


_GREEN = "C6EFCE"
_RED = "FFC7CE"
_YELLOW = "FFEB9C"
_BLUE = "BDD7EE"
_HEADER_BG = "4472C4"
_HEADER_FG = "FFFFFF"


def write_styled_excel(
    sheets: dict[str, pd.DataFrame],
    output_path: str,
    title: Optional[str] = None,
) -> str:
    """Write multiple DataFrames as styled sheets in one Excel workbook.

    Args:
        sheets: Mapping of sheet_name -> DataFrame.
        output_path: Destination .xlsx path.
        title: Optional title row written above the header on the first sheet.

    Returns:
        Absolute path to the written file.
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            start_row = 0
            if title and sheet_name == next(iter(sheets)):
                start_row = 1
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=start_row)
            ws = writer.sheets[sheet_name]

            if title and sheet_name == next(iter(sheets)):
                ws.cell(1, 1, title).font = Font(bold=True, size=14)

            header_row = start_row + 1
            for col_idx, _ in enumerate(df.columns, 1):
                cell = ws.cell(header_row, col_idx)
                cell.fill = PatternFill("solid", fgColor=_HEADER_BG)
                cell.font = Font(bold=True, color=_HEADER_FG)
                cell.alignment = Alignment(horizontal="center")

            for col_idx, col in enumerate(df.columns, 1):
                max_len = max(
                    (len(str(v)) for v in df[col]),
                    default=0,
                )
                width = min(max(max_len, len(str(col))) + 2, 50)
                ws.column_dimensions[get_column_letter(col_idx)].width = width

    return str(Path(output_path).resolve())


def color_cells(ws, row: int, col: int, color: str) -> None:
    """Fill a single cell with a solid background color."""
    ws.cell(row, col).fill = PatternFill("solid", fgColor=color)
