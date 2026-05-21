"""Tool 10: Null Checker — comprehensive null/empty value analysis."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import pandas as pd


@dataclass
class NullReport:
    column_stats: dict[str, dict]
    row_flags: list[dict]
    health_score_pct: float
    critical_rows: list[int]
    output_path: Optional[str] = None

    def export_report(self, output_path: str) -> str:
        """Write a two-sheet Excel report.

        Sheet 1: Column summary with null percentages.
        Sheet 2: Row-level view with flagged mandatory nulls.

        Args:
            output_path: Destination .xlsx path.

        Returns:
            Absolute path to the written file.
        """
        from openpyxl.styles import PatternFill
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        col_df = pd.DataFrame([
            {"Column": col, **stats}
            for col, stats in self.column_stats.items()
        ])
        row_df = pd.DataFrame(self.row_flags) if self.row_flags else pd.DataFrame(
            columns=["Row", "Null_Count", "Null_Columns", "Has_Critical_Null"]
        )

        red = "FFC7CE"
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            col_df.to_excel(writer, sheet_name="Column_Summary", index=False)
            row_df.to_excel(writer, sheet_name="Row_Summary", index=False)

            ws = writer.sheets["Row_Summary"]
            fill = PatternFill("solid", fgColor=red)
            for row_idx in range(2, ws.max_row + 1):
                critical_cell = ws.cell(row_idx, row_df.columns.get_loc("Has_Critical_Null") + 1 if "Has_Critical_Null" in row_df.columns else 4)
                if str(critical_cell.value).upper() == "TRUE":
                    for col_idx in range(1, ws.max_column + 1):
                        ws.cell(row_idx, col_idx).fill = fill

        return str(Path(output_path).resolve())


def _is_null(val) -> bool:
    if pd.isna(val):
        return True
    if isinstance(val, str) and val.strip().lower() in {"", "null", "n/a", "na", "none", "nil", "-"}:
        return True
    return False


def _load(file_path: str) -> pd.DataFrame:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    if path.suffix.lower() == ".csv":
        return pd.read_csv(file_path, dtype=str)
    return pd.read_excel(file_path, dtype=str)


def check_nulls(
    file_path: str,
    mandatory_columns: Optional[list[str]] = None,
    output_path: Optional[str] = None,
) -> NullReport:
    """Comprehensive null/empty value analysis for a file.

    Args:
        file_path: Source Excel/CSV path.
        mandatory_columns: Columns that must not be null. Rows missing these
            are flagged as critical.
        output_path: If provided, write a two-sheet Excel report here.

    Returns:
        NullReport with per-column stats, row flags, health score, and critical rows.

    Raises:
        FileNotFoundError: If the source file does not exist.
    """
    df = _load(file_path)
    mandatory_columns = mandatory_columns or []

    if df.empty:
        return NullReport(
            column_stats={},
            row_flags=[],
            health_score_pct=0.0,
            critical_rows=[],
            output_path=output_path,
        )

    total_cells = df.shape[0] * df.shape[1]
    total_nulls = 0

    column_stats: dict[str, dict] = {}
    for col in df.columns:
        null_mask = df[col].apply(_is_null)
        null_count = int(null_mask.sum())
        null_pct = round(null_count / len(df) * 100, 1) if len(df) else 0.0
        sample_indices = df.index[null_mask].tolist()[:5]
        total_nulls += null_count
        column_stats[col] = {
            "null_count": null_count,
            "null_pct": null_pct,
            "total_rows": len(df),
            "sample_null_rows": sample_indices,
        }

    health_score = round((1 - total_nulls / total_cells) * 100, 2) if total_cells else 0.0

    row_flags: list[dict] = []
    critical_rows: list[int] = []

    for idx, row in df.iterrows():
        null_cols = [col for col in df.columns if _is_null(row[col])]
        has_critical = any(col in mandatory_columns for col in null_cols) if mandatory_columns else False
        if has_critical:
            critical_rows.append(int(idx))
        row_flags.append({
            "Row": int(idx),
            "Null_Count": len(null_cols),
            "Null_Columns": ", ".join(null_cols),
            "Has_Critical_Null": has_critical,
        })

    saved_path = None
    if output_path:
        report = NullReport(
            column_stats=column_stats,
            row_flags=row_flags,
            health_score_pct=health_score,
            critical_rows=critical_rows,
        )
        saved_path = report.export_report(output_path)

    return NullReport(
        column_stats=column_stats,
        row_flags=row_flags,
        health_score_pct=health_score,
        critical_rows=critical_rows,
        output_path=saved_path,
    )
